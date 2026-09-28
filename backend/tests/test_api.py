import io
import os
import tempfile
import zipfile

_tmp = tempfile.mkdtemp()
os.environ["KP_DATA_DIR"] = _tmp
os.environ["KP_FRONTEND_DIST"] = os.path.join(_tmp, "no-dist")

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.main import app  # noqa: E402
from app.russian import director_short, greeting, position_dative  # noqa: E402


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        r = c.post("/api/login", json={"username": "admin", "password": "admin"})
        assert r.status_code == 200
        yield c


def doc_text(content: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        return z.read("word/document.xml").decode()


SERVICE = {
    "name": "Первичная регистрация в соответствии с требованиями ЕАЭС",
    "category": "Регистрация",
    "keywords": "регистрация ЕАЭС РУ",
    "duration_text": "от 12 месяцев с даты предоставления материалов от Заказчика.",
    "price": None,
    "price_unit": "",
    "price_note": "Итоговая стоимость согласуется после аудита досье и выбранных блоков",
    "footnotes": "*При необходимости",
    "stages": [
        {"name": "Аудит имеющейся документации досье.", "duration": "10 р.д."},
        {"name": "Разработка ПУР", "duration": "30 р.д."},
    ],
}


def test_requires_login():
    with TestClient(app) as c:
        assert c.get("/api/services").status_code == 401


def test_russian_forms():
    assert director_short("Иванов Иван Иванович") == "Иванову И.И."
    assert director_short("Петрова Анна Сергеевна") == "Петровой А.С."
    assert greeting("Петрова Анна Сергеевна") == "Уважаемая Анна Сергеевна!"
    assert greeting("Иванов Иван Иванович") == "Уважаемый Иван Иванович!"
    assert position_dative("Генеральный директор") == "Генеральному директору"
    assert position_dative("Коммерческий директор") == "Коммерческому директору"
    assert position_dative("Директор по развитию") == "Директору по развитию"
    assert position_dative("Управляющий") == "Управляющему"


def test_full_flow(client):
    r = client.post("/api/services", json=SERVICE)
    assert r.status_code == 200, r.text
    sid = r.json()["id"]
    assert len(r.json()["stages"]) == 2

    r2 = client.post("/api/services", json={
        "name": "Разработка мастер-файла системы фармаконадзора (МФСФ)",
        "keywords": "МФСФ", "price": 150000, "price_unit": "единовременно",
        "duration_text": "20 рабочих дней с момента получения материалов от Заказчика."})
    mid = r2.json()["id"]

    # автоподсказки: по началу слова, по ключевому слову и с опечаткой
    assert client.get("/api/services/search", params={"q": "мфсф"}).json()[0]["id"] == mid
    assert client.get("/api/services/search", params={"q": "первичная рег"}).json()[0]["id"] == sid
    assert client.get("/api/services/search", params={"q": "регистрацыя"}).json()[0]["id"] == sid

    draft = client.post("/api/proposals/draft", json={
        "company": "ООО «Ромашка»", "director_full": "Иванов Иван Иванович",
        "director_position": "Генеральный директор", "service_ids": [sid, mid]}).json()
    assert draft["director_short"] == "Иванову И.И."
    assert draft["items"][0]["stages"][1]["name"] == "Разработка ПУР"

    # правка в предпросмотре: новая цена + обновить её в реестре
    draft["items"][1]["price"] = 175000
    r = client.post("/api/proposals", json={"data": draft, "update_registry_prices": [mid]})
    assert r.status_code == 200, r.text
    pid = r.json()["id"]
    assert client.get(f"/api/services/{mid}").json()["price"] == "175000.00"

    doc = client.get(f"/api/proposals/{pid}/docx")
    assert doc.status_code == 200
    xml = doc_text(doc.content)
    for s in ["ООО «Ромашка»", "Иванову И.И.", "Уважаемый Иван Иванович!",
              "175 000 / единовременно", "Разработка ПУР", "30 р.д.",
              "Итоговая стоимость согласуется", "*При необходимости",
              "Крашенинников"]:
        assert s.replace("«", "&#171;") in xml or s in xml, s
    assert "{{" not in xml and "{%" not in xml

    # история
    detail = client.get(f"/api/proposals/{pid}").json()
    assert detail["data"]["items"][1]["price"] == "175000"
    assert client.get("/api/proposals", params={"q": "ромашка"}).json()[0]["id"] == pid


def test_no_stages_block_when_disabled(client):
    sid = client.post("/api/services", json={**SERVICE, "name": "Услуга с этапами 2"}).json()["id"]
    draft = client.post("/api/proposals/draft", json={"service_ids": [sid]}).json()
    draft["items"][0]["show_stages"] = False
    xml = doc_text(client.post("/api/proposals/preview-docx", json=draft).content)
    assert "Аудит имеющейся документации" not in xml
    assert "Перечень оказываемых услуг" not in xml


def test_excel_roundtrip(client):
    content = client.get("/api/registry/export").content
    wb = load_workbook(io.BytesIO(content))
    ws = wb["Услуги"]
    header = [c.value for c in ws[1]]
    name_col = header.index("Название*")
    price_col = header.index("Цена")
    # поменяем цену у существующей и добавим новую услугу с этапами
    ws.cell(row=2, column=price_col + 1, value=99000)
    ws.append([None, "Новая услуга из Excel", "Импорт"] + [None] * (len(header) - 3))
    ws.cell(row=ws.max_row, column=price_col + 1, value="12 500,50")
    wb["Этапы"].append([None, "Новая услуга из Excel", 2, "Второй этап", "5 р.д."])
    wb["Этапы"].append([None, "Новая услуга из Excel", 1, "Первый этап", "3 р.д."])
    first_name = ws.cell(row=2, column=name_col + 1).value
    buf = io.BytesIO()
    wb.save(buf)

    r = client.post("/api/registry/import",
                    files={"file": ("r.xlsx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 200, r.text
    rep = r.json()
    assert rep["created"] == 1 and rep["errors"] == []

    items = client.get("/api/services").json()
    new = next(s for s in items if s["name"] == "Новая услуга из Excel")
    full = client.get(f"/api/services/{new['id']}").json()
    assert full["price"] == "12500.50"
    assert [s["name"] for s in full["stages"]] == ["Первый этап", "Второй этап"]
    first = next(s for s in items if s["name"] == first_name)
    assert first["price"] == "99000.00"


def test_portal_base_path(monkeypatch):
    monkeypatch.setenv("BASE_PATH", "/kp")
    from app.config import cookie_path, normalized_base_path
    from app.main import create_app

    assert normalized_base_path() == "/kp"
    assert cookie_path() == "/kp"
    with TestClient(create_app()) as c:
        assert c.get("/kp/api/health").json() == {"ok": True}
        assert c.get("/api/health").status_code == 404
        login = c.post("/kp/api/login", json={"username": "admin", "password": "admin"})
        assert login.status_code == 200, login.text
        assert c.get("/kp/api/me").status_code == 200
