"""Операции с реестром услуг: сохранение, поиск, импорт/экспорт Excel."""

from __future__ import annotations

import io
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rapidfuzz import fuzz
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from . import defaults as D
from .models import Service, Stage
from .schemas import ServiceIn

# ------------------------------------------------------------ сохранение

SERVICE_FIELDS = [f for f in ServiceIn.model_fields if f != "stages"]


def apply_service(svc: Service, data: ServiceIn, username: str = "") -> Service:
    for f in SERVICE_FIELDS:
        setattr(svc, f, getattr(data, f))
    svc.stages.clear()
    for i, st in enumerate(data.stages):
        if st.name.strip():
            svc.stages.append(Stage(position=i, name=st.name.strip(),
                                    duration=st.duration.strip()))
    svc.updated_by = username
    return svc


# ------------------------------------------------------------ поиск

def _norm(s: str) -> str:
    return (s or "").lower().replace("ё", "е")


def search_services(db: Session, q: str, limit: int = 10) -> list[Service]:
    """Нечёткий поиск: по названию, ключевым словам и категории.

    Реестр небольшой (сотни позиций), поэтому ранжируем в Python — так
    корректно работает регистр кириллицы и опечатки.
    """
    services = db.scalars(select(Service).where(Service.active.is_(True))
                          .options(selectinload(Service.stages))).all()
    q = _norm(q).strip()
    if not q:
        return sorted(services, key=lambda s: s.name.lower())[:limit]
    words = q.split()
    scored = []
    for s in services:
        name = _norm(s.name)
        hay = f"{name} {_norm(s.keywords)} {_norm(s.category)}"
        if all(w in hay for w in words):
            score = 100 + (20 if name.startswith(q) else 0) + (10 if q in name else 0)
        else:
            score = max(fuzz.partial_ratio(q, name),
                        fuzz.token_set_ratio(q, hay),
                        fuzz.partial_ratio(q, _norm(s.keywords)) if s.keywords else 0)
        if score >= 65:
            scored.append((score, s))
    scored.sort(key=lambda x: (-x[0], x[1].name.lower()))
    return [s for _, s in scored[:limit]]


# ------------------------------------------------------------ Excel

# (заголовок колонки, поле модели)
SERVICE_COLUMNS = [
    ("ID", "id"),
    ("Название*", "name"),
    ("Категория", "category"),
    ("Группа в таблице КП", "table_group"),
    ("Ключевые слова", "keywords"),
    ("Описание", "description"),
    ("Подпись срока", "duration_label"),
    ("Срок выполнения", "duration_text"),
    ("Цена", "price"),
    ("Цена текстом", "price_text"),
    ("Единица цены", "price_unit"),
    ("Примечание к цене", "price_note"),
    ("Заголовок раздела", "section_title"),
    ("Вводный текст", "intro_text"),
    ("Заголовок колонки цены", "price_header"),
    ("Заголовок детализации", "stages_title"),
    ("Сноски", "footnotes"),
    ("Доп. текст", "extra_text"),
    ("Активна", "active"),
]
STAGE_COLUMNS = [
    ("ID услуги", "service_id"),
    ("Название услуги", "service_name"),
    ("№", "position"),
    ("Этап*", "name"),
    ("Срок", "duration"),
]
SHEET_SERVICES = "Услуги"
SHEET_STAGES = "Этапы"

_HEADER_FILL = PatternFill("solid", fgColor="C5E0B3")


def _write_sheet(ws, columns, rows):
    ws.append([c[0] for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    for i, (title, _) in enumerate(columns, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = 45 if title in (
            "Название*", "Описание", "Этап*", "Название услуги", "Заголовок детализации") else 18
    ws.freeze_panes = "A2"


def export_registry(db: Session, empty: bool = False) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_SERVICES
    services = [] if empty else db.scalars(
        select(Service).options(selectinload(Service.stages)).order_by(Service.category, Service.name)).all()

    rows = []
    for s in services:
        row = []
        for _, f in SERVICE_COLUMNS:
            v = getattr(s, f)
            if f == "active":
                v = "да" if v else "нет"
            elif f == "price" and v is not None:
                v = float(v)
            row.append(v)
        rows.append(row)
    if empty:
        rows.append([None, "Разработка мастер-файла системы фармаконадзора (МФСФ)", "Фармаконадзор",
                     "", "МФСФ, мастер-файл", "", D.DURATION_LABEL,
                     "20 рабочих дней с момента получения материалов от Заказчика.",
                     150000, "", "единовременно", "", "", "", "", "", "", "", "да"])
    _write_sheet(ws, SERVICE_COLUMNS, rows)

    ws2 = wb.create_sheet(SHEET_STAGES)
    st_rows = []
    for s in services:
        for i, st in enumerate(s.stages, start=1):
            st_rows.append([s.id, s.name, i, st.name, st.duration])
    if empty:
        st_rows.append([None, "Разработка мастер-файла системы фармаконадзора (МФСФ)", 1,
                        "Аудит имеющейся документации", "10 р.д."])
    _write_sheet(ws2, STAGE_COLUMNS, st_rows)

    ws3 = wb.create_sheet("Инструкция")
    for line in [
        "Как заполнять реестр услуг",
        "",
        "Лист «Услуги»: одна строка = одна услуга. Обязательна только колонка «Название*».",
        "ID — оставьте пустым для новой услуги. Если ID указан, услуга с этим ID будет обновлена.",
        "Если ID пуст, но услуга с таким же названием уже есть — она будет обновлена.",
        "Цена — число без пробелов и «руб.». Если цены нет (согласуется отдельно) — оставьте пустой",
        "  или заполните «Цена текстом» (например «по запросу»).",
        "Единица цены: единовременно, ежемесячно, за одну страну и т.п.",
        "Описание: можно в несколько строк (Alt+Enter). **текст** — выделится жирным.",
        "Сноски / Доп. текст: каждая строка ячейки — отдельный абзац в КП.",
        "Активна: да/нет (неактивные не предлагаются при создании КП).",
        "",
        "Лист «Этапы»: детализация (этапы и сроки). Одна строка = один этап.",
        "Услугу указывайте по ID или по точному названию. Порядок — по колонке «№».",
        "Если у услуги есть строки на листе «Этапы», её этапы в системе полностью заменяются.",
    ]:
        ws3.append([line])
    ws3["A1"].font = Font(bold=True, size=13)
    ws3.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _read_rows(ws, columns):
    it = ws.iter_rows(values_only=True)
    try:
        header = [(_cell_str(h)).rstrip("*").lower() for h in next(it)]
    except StopIteration:
        return []
    index = {}
    for title, field in columns:
        key = title.rstrip("*").lower()
        if key in header:
            index[field] = header.index(key)
    rows = []
    for n, raw in enumerate(it, start=2):
        if raw is None or all(v is None or _cell_str(v) == "" for v in raw):
            continue
        rows.append((n, {f: (raw[i] if i < len(raw) else None) for f, i in index.items()}))
    return rows


def _parse_price(v) -> Decimal | None:
    if v is None or _cell_str(v) == "":
        return None
    s = _cell_str(v).replace(" ", "").replace(" ", "").replace(",", ".")
    s = s.lower().replace("руб.", "").replace("руб", "").replace("₽", "")
    try:
        return Decimal(s)
    except InvalidOperation:
        raise ValueError(f"не удалось распознать цену «{_cell_str(v)}»")


def import_registry(db: Session, content: bytes, username: str = "") -> dict:
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as e:  # noqa: BLE001
        raise ValueError(f"Не удалось открыть файл Excel: {e}")
    if SHEET_SERVICES not in wb.sheetnames:
        raise ValueError(f"В файле нет листа «{SHEET_SERVICES}». Скачайте шаблон реестра.")

    report = {"created": 0, "updated": 0, "stages_services": 0, "errors": []}
    existing = db.scalars(select(Service).options(selectinload(Service.stages))).all()
    by_id = {s.id: s for s in existing}
    by_name = {_norm(s.name).strip(): s for s in existing}

    for n, row in _read_rows(wb[SHEET_SERVICES], SERVICE_COLUMNS):
        name = _cell_str(row.get("name"))
        if not name:
            report["errors"].append(f"Услуги, строка {n}: не указано название")
            continue
        try:
            price = _parse_price(row.get("price"))
        except ValueError as e:
            report["errors"].append(f"Услуги, строка {n}: {e}")
            continue
        sid = row.get("id")
        svc = None
        if sid not in (None, ""):
            try:
                svc = by_id.get(int(float(sid)))
            except (TypeError, ValueError):
                pass
        svc = svc or by_name.get(_norm(name).strip())
        if svc is None:
            svc = Service(name=name)
            db.add(svc)
            report["created"] += 1
        else:
            report["updated"] += 1
        svc.name = name
        svc.price = price
        for _, f in SERVICE_COLUMNS:
            if f in ("id", "name", "price") or f not in row:
                continue
            v = row[f]
            if f == "active":
                svc.active = _cell_str(v).lower() not in ("нет", "no", "0", "false", "н")
            elif f == "duration_label" and _cell_str(v) == "":
                svc.duration_label = D.DURATION_LABEL
            else:
                setattr(svc, f, _cell_str(v))
        svc.updated_by = username
        db.flush()
        by_id[svc.id] = svc
        by_name[_norm(name).strip()] = svc

    if SHEET_STAGES in wb.sheetnames:
        grouped: dict[int, list[tuple[float, str, str]]] = {}
        for n, row in _read_rows(wb[SHEET_STAGES], STAGE_COLUMNS):
            stage_name = _cell_str(row.get("name"))
            if not stage_name:
                report["errors"].append(f"Этапы, строка {n}: не указан этап")
                continue
            svc = None
            if _cell_str(row.get("service_id")):
                try:
                    svc = by_id.get(int(float(row["service_id"])))
                except (TypeError, ValueError):
                    pass
            if svc is None and _cell_str(row.get("service_name")):
                svc = by_name.get(_norm(_cell_str(row["service_name"])))
            if svc is None:
                report["errors"].append(f"Этапы, строка {n}: услуга не найдена")
                continue
            try:
                pos = float(row.get("position") or n)
            except (TypeError, ValueError):
                pos = n
            grouped.setdefault(svc.id, []).append((pos, stage_name, _cell_str(row.get("duration"))))
        for sid, stages in grouped.items():
            svc = by_id[sid]
            svc.stages.clear()
            db.flush()
            for i, (_, name, dur) in enumerate(sorted(stages, key=lambda x: x[0])):
                svc.stages.append(Stage(position=i, name=name, duration=dur))
            report["stages_services"] += 1

    db.commit()
    return report
